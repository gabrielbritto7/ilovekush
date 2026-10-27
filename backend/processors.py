from __future__ import annotations

import difflib
import hashlib
import threading
import urllib.request
import html
import io
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import fitz
from PIL import Image, ImageEnhance, ImageFilter, ImageOps

ROOT = Path(__file__).resolve().parents[1]
FONT = ROOT / 'fonts/DejaVuSans.ttf'
if not FONT.exists():
    FONT = ROOT / 'backend/fonts/DejaVuSans.ttf'
MODEL_DIR = Path(os.getenv('KUSH_MODEL_DIR', str(ROOT / '.data/models')))
MAX_PAGES = int(os.getenv('KUSH_MAX_PAGES', '500'))


class ProcessingError(Exception):
    pass


VALIDATOR_SHA512 = "abd403e96c26116bc0082609f3813b7cc3cd53fda068ff8abc76c8540a4d6ab638cb0ab525f04aa9c02b1406607594b18abae60848948dfb96b6ae909408ff6d"
_validator_lock = threading.Lock()


def validator_jar() -> Path:
    override=os.getenv('KUSH_VERAPDF_JAR')
    dest=Path(override) if override else ROOT/'.tools/verapdf/bin/cli-1.30.3.jar'
    if not dest.exists() and os.getenv('VERCEL') and not override:
        dest=Path(tempfile.gettempdir())/'ilovekush-verapdf-1.30.3.jar'
    with _validator_lock:
        if dest.exists():
            if hashlib.sha512(dest.read_bytes()).hexdigest()!=VALIDATOR_SHA512:
                raise ProcessingError('A verificação de integridade do validador PDF/A falhou.')
            return dest
        if not os.getenv('VERCEL') and not override:
            raise ProcessingError('Instale o veraPDF com scripts/install-verapdf.py.')
        url='https://repo.maven.apache.org/maven2/org/verapdf/apps/cli/1.30.3/cli-1.30.3.jar'
        try:
            data=urllib.request.urlopen(url,timeout=60).read(25*1024*1024)
            if hashlib.sha512(data).hexdigest()!=VALIDATOR_SHA512:
                raise ProcessingError('A verificação de integridade do validador PDF/A falhou.')
            dest.parent.mkdir(parents=True,exist_ok=True)
            fd,temporary=tempfile.mkstemp(prefix='kush-validator-',dir=dest.parent)
            with os.fdopen(fd,'wb') as stream:stream.write(data)
            os.replace(temporary,dest)
        except ProcessingError:raise
        except Exception:
            raise ProcessingError('Não foi possível preparar o validador PDF/A. Tente novamente em alguns instantes.')
    return dest


def command(args: list[str], timeout: int = 180) -> subprocess.CompletedProcess:
    try:
        result = subprocess.run(args, capture_output=True, text=True, timeout=timeout, check=False)
    except FileNotFoundError:
        raise ProcessingError('O motor necessário não está instalado neste servidor. Use o processamento no navegador quando disponível.')
    except subprocess.TimeoutExpired:
        raise ProcessingError('O documento excedeu o tempo de processamento. Tente um arquivo menor.')
    if result.returncode:
        raise ProcessingError('O motor não conseguiu processar este documento. Confira se o arquivo e as opções estão corretos.')
    return result


def number(options: dict, key: str, default: float, low: float, high: float) -> float:
    try:
        value = float(options.get(key, default))
    except (TypeError, ValueError):
        raise ProcessingError(f'O valor de {key} deve ser numérico.')
    if not math.isfinite(value) or not low <= value <= high:
        raise ProcessingError(f'O valor de {key} deve estar entre {low:g} e {high:g}.')
    return value


def pages(selection: str | list | None, count: int, *, unique: bool = True) -> list[int]:
    if not selection or str(selection).strip().lower() in ('all', 'todas', '*'):
        return list(range(count))
    result: list[int] = []
    for token in str(selection).replace(' ', '').split(','):
        if re.fullmatch(r'\d+-\d+', token):
            start, end = map(int, token.split('-'))
            if end < start:
                raise ProcessingError('O fim do intervalo deve ser maior ou igual ao início.')
            values = range(start, end + 1)
        elif token.isdigit():
            values = [int(token)]
        else:
            raise ProcessingError('Informe páginas como 1,3,5-8 ou deixe vazio para todas.')
        for value in values:
            if not 1 <= value <= count:
                raise ProcessingError(f'A página {value} não existe neste documento ({count} páginas).')
            if not unique or value - 1 not in result:
                result.append(value - 1)
    if not result:
        raise ProcessingError('Selecione pelo menos uma página.')
    return result


def open_pdf(path: Path, options: dict | None = None) -> fitz.Document:
    try:
        doc = fitz.open(path)
        if not doc.is_pdf:
            doc.close()
            raise ProcessingError('Envie um arquivo PDF válido.')
        if doc.needs_pass and not doc.authenticate(str((options or {}).get('password', ''))):
            doc.close()
            raise ProcessingError('Este PDF precisa da senha de abertura. Informe a senha correta nas opções.')
        if len(doc) > MAX_PAGES:
            doc.close()
            raise ProcessingError(f'Este servidor aceita até {MAX_PAGES} páginas por documento.')
        if not len(doc):
            doc.close()
            raise ProcessingError('O PDF não contém páginas.')
        return doc
    except ProcessingError:
        raise
    except Exception:
        raise ProcessingError('Não foi possível abrir este PDF. Verifique o arquivo ou tente Reparar PDF.')


def save(doc: fitz.Document, dest: Path, **kwargs) -> Path:
    doc.save(dest, garbage=4, deflate=True, clean=True, **kwargs)
    return dest


def color(value: str = '#204e3e') -> tuple[float, float, float]:
    if not re.fullmatch(r'#[\da-fA-F]{6}', str(value)):
        raise ProcessingError('Escolha uma cor válida.')
    return tuple(int(value[i:i + 2], 16) / 255 for i in (1, 3, 5))


def insert_text(page: fitz.Page, rect: fitz.Rect, text: str, fontsize: float = 14, ink='#204e3e', align=0, opacity=1.0):
    family = 'KushFont' if FONT.exists() else 'helv'
    fontfile = str(FONT) if FONT.exists() else None
    if rect.is_empty or rect.is_infinite:
        raise ProcessingError('A posição ou o tamanho do texto está fora da página.')
    result = page.insert_textbox(rect, text, fontsize=fontsize, fontname=family, fontfile=fontfile,
                                 color=color(ink), align=align, overlay=True, fill_opacity=opacity)
    if result < 0:
        raise ProcessingError('O texto não cabe na área escolhida. Diminua a fonte ou aumente a altura da área.')


def text_pdf(text: str, dest: Path, title='ILOVEKUSH') -> Path:
    doc=fitz.open()
    family='KushFont' if FONT.exists() else 'helv'
    face=fitz.Font(fontfile=str(FONT)) if FONT.exists() else fitz.Font('helv')
    page=doc.new_page();y=52
    def draw(value,size=11):
        nonlocal page,y
        if y>790:page=doc.new_page();y=52
        if FONT.exists():page.insert_font(fontname=family,fontfile=str(FONT))
        page.insert_text((44,y),value,fontname=family,fontsize=size,color=color())
        y+=size*1.55
    line=''
    for word in title.split():
        if face.text_length((line+' '+word).strip(),fontsize=20)>505 and line:draw(line,20);line=word
        else:line=(line+' '+word).strip()
    draw(line,20);y+=15
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            candidate=(line+' '+word).strip()
            if face.text_length(candidate,fontsize=11)>505 and line:draw(line);line=word
            else:line=candidate
        draw(line);y+=4
    doc.set_metadata({'title':title,'author':'ILOVEKUSH'})
    return save(doc,dest)


def inspect(path: Path, options: dict) -> dict:
    with open_pdf(path, options) as doc:
        fields = []
        for i, page in enumerate(doc):
            for widget in page.widgets() or []:
                fields.append(dict(name=widget.field_name, value=widget.field_value, type=widget.field_type_string,
                                   page=i + 1, choices=widget.choice_values or []))
        return dict(pages=len(doc), metadata=doc.metadata, fields=fields,
                    sizes=[dict(width=p.rect.width, height=p.rect.height) for p in doc])


def process(tool: str, files: list[Path], options: dict, out: Path, assets: dict[str, Path] | None = None) -> list[Path]:
    assets = assets or {}
    out.mkdir(exist_ok=True, parents=True)
    results: list[Path] = []
    if tool == 'merge':
        if len(files) < 2:
            raise ProcessingError('Selecione pelo menos dois PDFs para juntar.')
        merged = fitz.open()
        for file in files:
            with open_pdf(file, options) as source:
                merged.insert_pdf(source, widgets=True)
        return [save(merged, out / 'ilovekush-juntos.pdf')]
    if tool in ('jpg-pdf', 'scan'):
        doc = fitz.open()
        for file in files:
            with Image.open(file) as original:
                for frame in range(getattr(original, 'n_frames', 1)):
                    original.seek(frame)
                    image = ImageOps.exif_transpose(original).convert('RGB')
                    if image.width * image.height > 60_000_000:
                        raise ProcessingError('A imagem é grande demais. Use até 60 megapixels.')
                    if tool == 'scan':
                        image = ImageOps.autocontrast(image)
                        if options.get('scan_mode', 'color') == 'gray':
                            image = ImageOps.grayscale(image).convert('RGB')
                        elif options.get('scan_mode') == 'bw':
                            image = ImageOps.grayscale(image).point(lambda x: 255 if x > 160 else 0).convert('RGB')
                    size = fitz.paper_rect('a4') if options.get('paper') == 'a4' else fitz.Rect(0, 0, image.width * .75, image.height * .75)
                    if options.get('orientation') == 'landscape' and size.height > size.width:
                        size = fitz.Rect(0, 0, size.height, size.width)
                    margin = number(options, 'margin', 12, 0, min(size.width, size.height) / 4)
                    page = doc.new_page(width=size.width, height=size.height)
                    stream = io.BytesIO();image.save(stream, format='JPEG', quality=92)
                    page.insert_image(page.rect + (margin, margin, -margin, -margin), stream=stream.getvalue(), keep_proportion=True)
        return [save(doc, out / 'ilovekush-imagens.pdf')]
    if tool in ('word-pdf', 'excel-pdf', 'powerpoint-pdf'):
        executable = shutil.which('soffice')
        if not executable:
            raise ProcessingError('Use o conversor Office no navegador. O LibreOffice não está instalado neste servidor.')
        profile = out / 'office-profile'
        for i, file in enumerate(files):
            target = out / f'office-{i}'; target.mkdir()
            command([executable, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(target), str(file)])
            found = list(target.glob('*.pdf'))
            if not found:
                raise ProcessingError('O LibreOffice não conseguiu converter este arquivo. Confira o formato e a senha do documento.')
            dest = out / f'{file.stem}.pdf';shutil.copyfile(found[0], dest);results.append(dest)
        return results
    if tool == 'html-pdf':
        from weasyprint import HTML
        def restricted_fetcher(url, *args, **kwargs):
            if str(url).startswith('data:image/'):
                from weasyprint import default_url_fetcher
                return default_url_fetcher(url)
            raise ValueError('Recursos externos não são permitidos.')
        text = str(options.get('html') or (files[0].read_text(errors='replace') if files else ''))
        if not text.strip():
            raise ProcessingError('Cole o HTML ou envie um arquivo HTML.')
        dest = out / 'ilovekush-html.pdf'
        HTML(string=text, url_fetcher=restricted_fetcher).write_pdf(dest)
        return [dest]
    if tool == 'compare':
        if len(files) != 2:
            raise ProcessingError('Selecione exatamente dois PDFs para comparar.')
        with open_pdf(files[0], options) as left, open_pdf(files[1], options) as right:
            ltext='\n'.join(p.get_text() for p in left);rtext='\n'.join(p.get_text() for p in right)
            diff='\n'.join(difflib.unified_diff(ltext.splitlines(), rtext.splitlines(), fromfile=files[0].name, tofile=files[1].name, lineterm=''))
            (out/'diferencas.txt').write_text(diff or 'Nenhuma diferença de texto encontrada. Confira também a comparação visual.', encoding='utf-8')
            report = fitz.open()
            for index in range(max(len(left), len(right))):
                canvas=Image.new('RGB',(1240,900),'#f7f6f1')
                rendered=[]
                for d in [left,right]:
                    if index < len(d):
                        pix=d[index].get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
                        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
                        im.thumbnail((590,850));rendered.append(im)
                    else:rendered.append(Image.new('RGB',(590,850),'white'))
                from PIL import ImageChops
                from PIL import ImageDraw
                same=Image.new('RGB',(590,850),'white');same.paste(rendered[0],(0,0))
                other=Image.new('RGB',(590,850),'white');other.paste(rendered[1],(0,0))
                mask=ImageChops.difference(same,other).convert('L').point(lambda p:255 if p>40 else 0)
                if mask.getbbox():
                    overlay=Image.new('RGB',other.size,'#f29b9b');other=Image.composite(Image.blend(other,overlay,.5),other,mask)
                canvas.paste(same,(20,35));canvas.paste(other,(630,35))
                draw=ImageDraw.Draw(canvas);draw.text((20,12),f'Original - {index+1}',fill='#204e3e');draw.text((630,12),f'Comparado - {index+1}',fill='#204e3e')
                stream=io.BytesIO();canvas.save(stream,format='PNG')
                page=report.new_page(width=842,height=611);page.insert_image(page.rect,stream=stream.getvalue())
            save(report,out/'comparacao-visual.pdf')
        return [out/'diferencas.txt',out/'comparacao-visual.pdf']

    for index, file in enumerate(files):
        stem = file.stem
        dest = out / f'{stem}-{tool}.pdf'
        with open_pdf(file, options) as doc:
            selection = pages(options.get('pages'), len(doc))
            if tool == 'split':
                mode=options.get('split_mode','pages')
                if mode=='ranges':
                    groups=[pages(part.strip(),len(doc)) for part in str(options.get('ranges','')).split(';') if part.strip()]
                    if not groups:raise ProcessingError('Informe intervalos separados por ponto e vírgula: 1-3;4-6.')
                elif mode=='fixed':
                    count=int(number(options,'group_size',1,1,len(doc)));groups=[list(range(i,min(i+count,len(doc)))) for i in range(0,len(doc),count)]
                elif mode=='extract':groups=[selection]
                elif mode=='remove':
                    remaining=[i for i in range(len(doc)) if i not in selection]
                    if not remaining:raise ProcessingError('A remoção deixaria o documento sem páginas.')
                    groups=[remaining]
                else:groups=[[i] for i in selection]
                for j,group in enumerate(groups):
                    part=fitz.open()
                    for p in group:part.insert_pdf(doc,from_page=p,to_page=p)
                    results.append(save(part,out/f'{stem}-parte-{j+1}.pdf'))
            elif tool == 'organize':
                ordered=pages(options.get('order') or options.get('pages'),len(doc),unique=False)
                part=fitz.open()
                for p in ordered:part.insert_pdf(doc,from_page=p,to_page=p,widgets=True)
                results.append(save(part,dest))
            elif tool == 'rotate':
                angle=int(number(options,'angle',90,0,270))
                if angle not in [0,90,180,270]:raise ProcessingError('Escolha 0°, 90°, 180° ou 270°.')
                for p in selection:doc[p].set_rotation((doc[p].rotation+angle)%360)
                results.append(save(doc,dest))
            elif tool == 'compress':
                if shutil.which('gs'):
                    level={'low':'/prepress','recommended':'/ebook','extreme':'/screen'}.get(options.get('level','recommended'),'/ebook')
                    source=out/f'clear-{index}.pdf';doc.save(source,encryption=fitz.PDF_ENCRYPT_NONE)
                    command(['gs','-sDEVICE=pdfwrite','-dCompatibilityLevel=1.7',f'-dPDFSETTINGS={level}','-dNOPAUSE','-dBATCH','-dSAFER',f'-sOutputFile={dest}',str(source)])
                else:
                    doc.save(dest,garbage=4,deflate=True,clean=True,use_objstms=1)
                if dest.stat().st_size>=file.stat().st_size and not doc.needs_pass:
                    shutil.copyfile(file,dest)
                results.append(dest)
            elif tool == 'protect':
                password=str(options.get('new_password',''))
                if len(password)<4:raise ProcessingError('Use uma senha de pelo menos 4 caracteres.')
                results.append(save(doc,dest,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw=password,user_pw=password,permissions=fitz.PDF_PERM_PRINT|fitz.PDF_PERM_COPY))
            elif tool in ('unlock','repair'):
                results.append(save(doc,dest,encryption=fitz.PDF_ENCRYPT_NONE))
            elif tool == 'crop':
                margin=number(options,'crop_margin',24,0,200)
                for p in selection:
                    page=doc[p];rect=page.cropbox
                    if margin*2>=min(rect.width,rect.height):raise ProcessingError('As margens são grandes demais para esta página.')
                    page.set_cropbox(rect+(margin,margin,-margin,-margin))
                results.append(save(doc,dest))
            elif tool in ('edit','watermark','pagenumber','sign'):
                text=str(options.get('text',''))
                for position,p in enumerate(selection):
                    page=doc[p];w,h=page.rect.width,page.rect.height
                    fontsize=number(options,'font_size',24 if tool=='watermark' else 16,6,180)
                    x=number(options,'x',10,0,95)/100*w
                    y=number(options,'y',80 if tool=='sign' else 12,0,95)/100*h
                    if tool=='pagenumber':
                        text=str(options.get('format','{n} / {total}')).replace('{n}',str(int(number(options,'start',1,0,999999))+p)).replace('{total}',str(len(doc)))
                        top=options.get('position')=='top';rect=fitz.Rect(24,20 if top else h-40,w-24,50 if top else h-10)
                        insert_text(page,rect,text,fontsize=12,ink=str(options.get('color','#204e3e')),align=1)
                    elif tool=='watermark':
                        opacity=number(options,'opacity',30,5,100)/100
                        if 'image' in assets:
                            page.insert_image(fitz.Rect(w*.2,h*.35,w*.8,h*.65),filename=str(assets['image']),keep_proportion=True,overlay=options.get('layer','above')=='above')
                        else:
                            if not text.strip():raise ProcessingError('Escreva a marca d’água ou envie uma imagem.')
                            insert_text(page,fitz.Rect(20,h*.45,w-20,h*.7),text,fontsize,ink=str(options.get('color','#204e3e')),align=1,opacity=opacity)
                    elif tool=='sign' and 'signature' in assets:
                        page.insert_image(fitz.Rect(x,y,min(w-10,x+w*.35),min(h-10,y+90)),filename=str(assets['signature']),keep_proportion=True)
                    else:
                        if tool=='edit' and options.get('find'):
                            hits=page.search_for(str(options['find']))
                            for hit in hits:page.add_redact_annot(hit,fill=(1,1,1))
                            page.apply_redactions()
                            for hit in hits:insert_text(page,fitz.Rect(hit.x0,hit.y0,min(w,hit.x0+max(hit.width,200)),min(h,hit.y0+fontsize*3)),text,fontsize,ink=str(options.get('color','#204e3e')))
                        elif text.strip():
                            insert_text(page,fitz.Rect(x,y,w-20,min(h-10,y+number(options,'height',100,20,500))),text,fontsize,ink=str(options.get('color','#204e3e')))
                        if tool=='edit' and 'image' in assets:
                            page.insert_image(fitz.Rect(x,y,min(w-10,x+200),min(h-10,y+150)),filename=str(assets['image']),keep_proportion=True)
                        if not text.strip() and not assets:raise ProcessingError('Adicione um texto, uma assinatura ou uma imagem.')
                results.append(save(doc,dest))
            elif tool == 'redact':
                needle=str(options.get('redact_text','')).strip()
                boxes=options.get('boxes',[])
                if not needle and not boxes:raise ProcessingError('Informe o texto ou marque uma região para remover.')
                hits=0
                for p in selection:
                    page=doc[p]
                    for word in [x.strip() for x in needle.split('\n') if x.strip()]:
                        for rect in page.search_for(word):page.add_redact_annot(rect,fill=(0,0,0));hits+=1
                    for box in boxes:
                        if int(box.get('page',p+1))!=p+1:continue
                        vals=[float(box[k]) for k in ['x','y','width','height']]
                        if any(not math.isfinite(v) or v<0 or v>100 for v in vals):raise ProcessingError('A região de censura deve estar dentro da página.')
                        x,y,bw,bh=vals;r=page.rect
                        page.add_redact_annot(fitz.Rect(r.width*x/100,r.height*y/100,min(r.width,r.width*(x+bw)/100),min(r.height,r.height*(y+bh)/100)),fill=(0,0,0));hits+=1
                    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,text=fitz.PDF_REDACT_TEXT_REMOVE)
                if not hits:raise ProcessingError('O texto não foi encontrado. Em um PDF digitalizado, selecione uma área para censurar.')
                doc.set_metadata({});doc.del_xml_metadata()
                for name in doc.embfile_names():doc.embfile_del(name)
                results.append(save(doc,dest))
            elif tool == 'pdf-jpg':
                if options.get('image_mode')=='extract':
                    seen=set()
                    for p in selection:
                        for item in doc[p].get_images(full=True):
                            if item[0] in seen:continue
                            seen.add(item[0]);image=doc.extract_image(item[0]);path=out/f'{stem}-imagem-{len(seen)}.{image["ext"]}';path.write_bytes(image['image']);results.append(path)
                    if not seen:raise ProcessingError('Não há imagens incorporadas. Use o modo Converter páginas para renderizar o documento.')
                else:
                    dpi=number(options,'dpi',150,72,300)
                    for p in selection:
                        pix=doc[p].get_pixmap(dpi=int(dpi),alpha=False)
                        path=out/f'{stem}-pagina-{p+1}.jpg';pix.save(path,jpg_quality=90);results.append(path)
            elif tool in ('extract','markdown','summarize','smart-split','translate'):
                texts=[p.get_text(sort=True) for p in doc];full='\n\n'.join(texts)
                if not full.strip():raise ProcessingError('Este documento não contém texto pesquisável. Execute OCR PDF primeiro.')
                if tool=='extract':
                    path=out/f'{stem}.{"json" if options.get("detailed") else "txt"}'
                    if options.get('detailed'):
                        value=[dict(page=i+1,words=p.get_text('words')) for i,p in enumerate(doc)];path.write_text(json.dumps(value,ensure_ascii=False,indent=2))
                    else:path.write_text(full,encoding='utf-8')
                    results.append(path)
                elif tool=='markdown':
                    lines=[]
                    for p in doc:
                        for block in p.get_text('dict',sort=True)['blocks']:
                            for line in block.get('lines',[]):
                                text=''.join(span['text'] for span in line['spans']);size=max((span['size'] for span in line['spans']),default=10)
                                lines.append(('## ' if size>=17 else '')+text)
                        lines.append('\n---\n')
                    path=out/f'{stem}.md';path.write_text('\n\n'.join(lines),encoding='utf-8');results.append(path)
                elif tool=='summarize':
                    sentences=[s.strip() for s in re.split(r'(?<=[.!?])\s+|\n\n+',full) if len(s.strip())>25]
                    if not sentences:sentences=[full.strip()]
                    stop=set('de da do das dos a o as os e em um uma para por com que é se não no na the of and is to in for on it a an'.split())
                    tokens=lambda text:re.findall(r'\b[^\W\d_]{3,}\b',text.lower())
                    freq=Counter(t for s in sentences for t in tokens(s) if t not in stop)
                    ranked=sorted(range(len(sentences)),key=lambda i:sum(freq[t] for t in tokens(sentences[i]) if t not in stop)/max(1,len(tokens(sentences[i]))) ,reverse=True)
                    count=min(len(sentences),int(number(options,'sentences',6,1,40)))
                    summary='\n\n'.join(sentences[i] for i in sorted(ranked[:count]))
                    if options.get('output')=='txt':path=out/f'{stem}-resumo.txt';path.write_text(summary,encoding='utf-8');results.append(path)
                    else:results.append(text_pdf(summary,out/f'{stem}-resumo.pdf','Resumo • '+stem))
                elif tool=='smart-split':
                    keyword=str(options.get('keyword','')).strip()
                    if not keyword:raise ProcessingError('Informe a palavra ou o título que marca o início de cada seção.')
                    boundaries=[0]+[i for i,text in enumerate(texts) if i>0 and keyword.casefold() in text.casefold()]+[len(doc)]
                    if len(boundaries)==2:raise ProcessingError('Não foram encontradas novas seções com esse texto. Confira a palavra-chave.')
                    for i,(start,end) in enumerate(zip(boundaries,boundaries[1:])):
                        part=fitz.open();part.insert_pdf(doc,from_page=start,to_page=end-1);results.append(save(part,out/f'{stem}-secao-{i+1}.pdf'))
                else:
                    from_lang=str(options.get('from','en'));to_lang=str(options.get('to','pt'))
                    model=MODEL_DIR/f'{from_lang}_{to_lang}'
                    if not (model/'model').is_dir():raise ProcessingError('O modelo desse par de idiomas ainda não está instalado. Instale-o conforme o README ou escolha outro par disponível.')
                    import ctranslate2
                    import sentencepiece as spm
                    sp=spm.SentencePieceProcessor(model_file=str(model/'sentencepiece.model'))
                    translator=ctranslate2.Translator(str(model/'model'),device='cpu',inter_threads=1,intra_threads=2)
                    chunks=[x.strip() for x in re.split(r'(?<=[.!?])\s+|\n+',full) if x.strip()]
                    chunks=[chunk[i:i+1200] for chunk in chunks for i in range(0,len(chunk),1200)]
                    translated=translator.translate_batch([sp.encode(x,out_type=str) for x in chunks],max_batch_size=8,beam_size=3)
                    text='\n\n'.join(sp.decode(x.hypotheses[0]) for x in translated)
                    results.append(text_pdf(text,out/f'{stem}-traduzido.pdf','Tradução • '+stem))
            elif tool=='pdf-word':
                target=out/f'{stem}.docx'
                if not os.getenv('VERCEL'):
                    from pdf2docx import Converter
                    conv=Converter(str(file),password=options.get('password'))
                    try:conv.convert(str(target),multi_processing=False)
                    finally:conv.close()
                else:
                    from docx import Document
                    word=Document()
                    for i,page in enumerate(doc):
                        for text in page.get_text(sort=True).split('\n'):word.add_paragraph(text)
                        if i<len(doc)-1:word.add_page_break()
                    word.save(target)
                results.append(target)
            elif tool=='pdf-excel':
                from openpyxl import Workbook
                wb=Workbook();wb.remove(wb.active)
                for i,page in enumerate(doc):
                    ws=wb.create_sheet(f'Página {i+1}')
                    tables=page.find_tables().tables
                    if tables:
                        for table in tables:
                            for row in table.extract():ws.append([str(value or '') for value in row])
                            ws.append([])
                    else:
                        for line in page.get_text(sort=True).splitlines():ws.append([line])
                    for row in ws:
                        for cell in row:
                            if isinstance(cell.value,str):cell.data_type='s'
                    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=30
                target=out/f'{stem}.xlsx';wb.save(target);results.append(target)
            elif tool=='pdf-powerpoint':
                from pptx import Presentation
                from pptx.util import Inches
                prs=Presentation();prs.slide_width=Inches(11.7);prs.slide_height=Inches(8.3)
                for page in doc:
                    image=page.get_pixmap(dpi=120,alpha=False);stream=io.BytesIO(image.tobytes('png'))
                    slide=prs.slides.add_slide(prs.slide_layouts[6]);ratio=min(prs.slide_width/image.width,prs.slide_height/image.height)
                    width=int(image.width*ratio);height=int(image.height*ratio)
                    slide.shapes.add_picture(stream,(prs.slide_width-width)//2,(prs.slide_height-height)//2,width=width,height=height)
                target=out/f'{stem}.pptx';prs.save(target);results.append(target)
            elif tool=='ocr':
                import pytesseract
                languages=str(options.get('language','por+eng'))
                if any(not re.fullmatch('[a-z_]+',x) for x in languages.split('+')):raise ProcessingError('Escolha um idioma de OCR válido.')
                tessdata=str(ROOT/'.data/tessdata')
                if (ROOT/'.data/tessdata/eng.traineddata').exists():os.environ['TESSDATA_PREFIX']=tessdata
                available=pytesseract.get_languages(config='')
                if any(x not in available for x in languages.split('+')):raise ProcessingError('O idioma selecionado não está instalado neste servidor. Use OCR no navegador ou instale os idiomas do Tesseract.')
                target=fitz.open()
                for page in doc:
                    pix=page.get_pixmap(dpi=160,alpha=False)
                    image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
                    pdf=pytesseract.image_to_pdf_or_hocr(image,extension='pdf',lang=languages,timeout=90)
                    with fitz.open(stream=pdf,filetype='pdf') as part:target.insert_pdf(part)
                results.append(save(target,dest))
            elif tool=='forms':
                fields=options.get('fields',{})
                if not isinstance(fields,dict):raise ProcessingError('Os campos do formulário são inválidos.')
                found=0
                for page in doc:
                    for widget in page.widgets() or []:
                        if widget.field_name in fields:
                            value=fields[widget.field_name]
                            widget.field_value=(widget.on_state() if value else 'Off') if widget.field_type==fitz.PDF_WIDGET_TYPE_CHECKBOX else str(value)
                            widget.update();found+=1
                if not found:raise ProcessingError('Nenhum campo foi preenchido. Abra um PDF com campos interativos e selecione os valores.')
                if options.get('flatten'):
                    flattened=fitz.open()
                    for page in doc:
                        pix=page.get_pixmap(dpi=150,alpha=False);p=flattened.new_page(width=page.rect.width,height=page.rect.height);p.insert_image(p.rect,stream=pix.tobytes('png'))
                    results.append(save(flattened,dest))
                else:results.append(save(doc,dest))
            elif tool=='pdfa':
                if not shutil.which('gs'):raise ProcessingError('O Ghostscript não está instalado. Use o processamento PDF/A no navegador.')
                icc=Path('/usr/share/color/icc/ghostscript/srgb.icc')
                if not icc.exists():icc=Path('/usr/share/color/icc/sRGB.icc')
                if not icc.exists():raise ProcessingError('O perfil de cores sRGB não foi encontrado.')
                clear=out/f'pdfa-source-{index}.pdf';doc.save(clear,encryption=fitz.PDF_ENCRYPT_NONE)
                ps=out/f'pdfa-{index}.ps';ps.write_text(f'''%!PS-Adobe-3.0\n/ICCProfile ({icc}) def\n[/_objdef {{icc_PDFA}} /type /stream /OBJ pdfmark\n[{{icc_PDFA}} << /N 3 >> /PUT pdfmark\n[{{icc_PDFA}} ICCProfile (r) file /PUT pdfmark\n[/_objdef {{OutputIntent_PDFA}} /type /dict /OBJ pdfmark\n[{{OutputIntent_PDFA}} << /Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {{icc_PDFA}} /OutputConditionIdentifier (sRGB) >> /PUT pdfmark\n[{{Catalog}} << /OutputIntents [{{OutputIntent_PDFA}}] >> /PUT pdfmark\n''')
                command(['gs','-dPDFA=2','-dBATCH','-dNOPAUSE','-dSAFER',f'--permit-file-read={icc}','-sDEVICE=pdfwrite','-sColorConversionStrategy=RGB','-dPDFACompatibilityPolicy=1',f'-sOutputFile={dest}',str(ps),str(clear)])
                results.append(dest)
            elif tool=='validate-pdfa':
                jar=validator_jar()
                java=shutil.which('java')
                if not java:
                    try:
                        from jdk4py import JAVA
                        java=str(JAVA)
                    except ImportError:pass
                if not java or not jar.exists():raise ProcessingError('O validador veraPDF ainda não está instalado neste servidor. Consulte o README para ativar a validação completa.')
                report=command([java,'-Xmx256m','-jar',str(jar),'--format','json',str(file)],timeout=120)
                path=out/f'{stem}-validacao.json';path.write_text(report.stdout,encoding='utf-8');results.append(path)
            elif tool=='certificate':
                if 'certificate' not in assets:raise ProcessingError('Envie seu certificado PFX ou P12.')
                from pyhanko.sign import signers
                from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
                signer=signers.SimpleSigner.load_pkcs12(str(assets['certificate']),passphrase=str(options.get('certificate_password','')).encode())
                if signer is None:raise ProcessingError('Não foi possível abrir o certificado. Confira o arquivo e a senha.')
                with file.open('rb') as source, dest.open('wb') as target:
                    writer=IncrementalPdfFileWriter(source)
                    if doc.needs_pass:writer.encrypt(str(options.get('password','')))
                    signers.sign_pdf(writer,signature_meta=signers.PdfSignatureMetadata(field_name='ILOVEKUSH',reason=str(options.get('reason','Assinatura digital'))),signer=signer,output=target)
                results.append(dest)
            else:
                raise ProcessingError('Ferramenta não encontrada.')
    if not results:
        raise ProcessingError('Nenhum resultado foi gerado. Confira os arquivos e as opções.')
    return results


def package(files: list[Path], dest: Path) -> Path:
    if len(files)==1:return files[0]
    used=set()
    with zipfile.ZipFile(dest,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for i,file in enumerate(files):
            name=file.name if file.name not in used else f'{i+1}-{file.name}';used.add(name);z.write(file,name)
    return dest
