import datetime
from pathlib import Path
import fitz
from PIL import Image

def build_documents(folder: Path):
    folder.mkdir(parents=True,exist_ok=True)
    pdf=fitz.open()
    for i in range(3):
        p=pdf.new_page()
        p.insert_text((60,80),f'Chapter {i+1}',fontsize=22)
        p.insert_text((60,120),f'ILOVEKUSH document page {i+1}. SECRET-123. This is a document about PDF processing.',fontsize=11)
        p.insert_text((60,150),'Open source tools make it possible to work with documents. Privacy is important.',fontsize=11)
        p.insert_text((60,180),'This sentence describes the processing of files. Each chapter contains useful information.',fontsize=11)
    pdf.save(folder/'sample.pdf');pdf.close()
    pdf=fitz.open();p=pdf.new_page();p.insert_text((60,80),'Another document, with different content.',fontsize=16);pdf.save(folder/'other.pdf');pdf.close()
    image=Image.new('RGB',(600,400),'#cfe7b9');image.save(folder/'image.png')
    from docx import Document
    word=Document();word.add_heading('ILOVEKUSH Word document',0);word.add_paragraph('A real editable document.');word.save(folder/'sample.docx')
    from openpyxl import Workbook
    wb=Workbook();wb.active.append(['Product','Value']);wb.active.append(['PDF',42]);wb.save(folder/'sample.xlsx')
    from pptx import Presentation
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text='ILOVEKUSH Presentation';prs.save(folder/'sample.pptx')
    pdf=fitz.open();p=pdf.new_page();widget=fitz.Widget();widget.field_name='name';widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT;widget.rect=fitz.Rect(60,60,300,95);p.add_widget(widget);pdf.save(folder/'form.pdf');pdf.close()
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes,serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'ILOVEKUSH Test Only')])
    cert=x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()).serial_number(x509.random_serial_number()).not_valid_before(datetime.datetime.now(datetime.timezone.utc)-datetime.timedelta(minutes=1)).not_valid_after(datetime.datetime.now(datetime.timezone.utc)+datetime.timedelta(days=1)).add_extension(x509.BasicConstraints(ca=False,path_length=None),critical=True).sign(key,hashes.SHA256())
    (folder/'test.p12').write_bytes(pkcs12.serialize_key_and_certificates(b'test',key,cert,None,serialization.BestAvailableEncryption(b'test-password')))
    (folder/'page.html').write_text('<h1>ILOVEKUSH HTML</h1><p>A real HTML document.</p>')
    pdf=fitz.open(folder/'sample.pdf');pix=pdf[0].get_pixmap(dpi=150);scan=fitz.open();p=scan.new_page(width=pdf[0].rect.width,height=pdf[0].rect.height);p.insert_image(p.rect,stream=pix.tobytes('png'));scan.save(folder/'scan.pdf');scan.close();pdf.close()
    return folder
